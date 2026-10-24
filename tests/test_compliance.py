import io
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from app import db
from app.models import Action, Assessment, Framework
from app.services import assessments as svc
from app.services import versioning
from app.services.spreadsheet import RequirementRow, SpreadsheetError, build_template, parse_requirements


def rows(*specs):
    """specs: (code, title) or (code, title, previous_code)."""
    return [RequirementRow(code=s[0], title=s[1], section="S", previous_code=s[2] if len(s) > 2 else "")
            for s in specs]


@pytest.fixture
def two_versions(app):
    fw = Framework(name="ISO")
    db.session.add(fw)
    v1, _ = versioning.create_version(fw, "2013", rows=rows(("A1", "Policy"), ("A2", "Roles"), ("A3", "Assets"),
                                                             ("A4", "Passwords")))
    versioning.publish(v1)
    # A1 unchanged, A2 reworded, A3 renumbered to B3, A4 removed, B5 new.
    v2, warnings = versioning.create_version(
        fw, "2022", base_version=v1,
        rows=rows(("A1", "Policy"), ("A2", "Roles and duties"), ("B3", "Assets", "A3"), ("B5", "Cloud")),
    )
    versioning.publish(v2)
    db.session.commit()
    assert warnings == []
    return fw, v1, v2


def by_code(version):
    return {r.code: r for r in version.requirements}


# --- spreadsheet -----------------------------------------------------------

def test_template_round_trip():
    parsed = parse_requirements("t.xlsx", build_template())
    assert [r.code for r in parsed] == ["1.1", "1.2", "2.1"]
    assert parsed[0].section == "1. Governance"


def test_parse_csv_with_aliases_and_semicolons():
    data = "ID;Name;Chapter;Old code\n5.10;Do X;Ch1;\n5.11;Do Y;Ch1;5.2\n".encode()
    parsed = parse_requirements("r.csv", io.BytesIO(data))
    assert [(r.code, r.title, r.section, r.previous_code) for r in parsed] == [
        ("5.10", "Do X", "Ch1", ""), ("5.11", "Do Y", "Ch1", "5.2")]


def test_parse_errors_are_reported():
    with pytest.raises(SpreadsheetError) as e:
        parse_requirements("r.csv", io.BytesIO(b"Code,Title\nA,One\nA,Two\n,Three\n"))
    assert len(e.value.errors) == 2
    with pytest.raises(SpreadsheetError):
        parse_requirements("r.csv", io.BytesIO(b"Code,Description\nA,One\n"))


# --- versioning --------------------------------------------------------------

def test_version_diff(two_versions):
    _, v1, v2 = two_versions
    reqs = by_code(v2)
    assert {c: r.change_type for c, r in reqs.items()} == {
        "A1": "unchanged", "A2": "modified", "B3": "modified", "B5": "new"}
    old = by_code(v1)
    assert reqs["B3"].lineage_id == old["A3"].lineage_id
    assert reqs["B5"].lineage_id == reqs["B5"].id
    assert [r.code for r in versioning.removed_requirements(v2)] == ["A4"]
    assert versioning.change_summary(v2) == {"new": 1, "modified": 2, "unchanged": 1, "removed": 1, "total": 4}


def test_clone_then_edit_draft(two_versions):
    fw, _, v2 = two_versions
    v3, _ = versioning.create_version(fw, "2025", base_version=v2)
    reqs = by_code(v3)
    assert all(r.change_type == "unchanged" for r in reqs.values())
    versioning.save_requirement(v3, {"code": "A1", "title": "Policy v2"}, reqs["A1"])
    versioning.save_requirement(v3, {"code": "C1", "title": "AI governance"})
    versioning.delete_requirement(reqs["B5"])
    db.session.flush()
    db.session.refresh(v3)
    summary = versioning.change_summary(v3)
    assert (summary["modified"], summary["new"], summary["removed"]) == (1, 1, 1)
    with pytest.raises(versioning.VersionError):
        versioning.save_requirement(v3, {"code": "A2", "title": "dup"})
    versioning.publish(v3)
    with pytest.raises(versioning.VersionError):
        versioning.save_requirement(v3, {"code": "Z", "title": "too late"})


def test_unknown_previous_code_warns(two_versions):
    fw, _, v2 = two_versions
    _, warnings = versioning.create_version(fw, "x", base_version=v2, rows=rows(("N1", "n", "NOPE")))
    assert "not found" in warnings[0]


# --- assessments -------------------------------------------------------------

def answer(a, mapping):
    for r in a.results:
        if r.requirement.code in mapping:
            svc.update_result(a, r, status=mapping[r.requirement.code])


def test_kpis_and_completion(two_versions):
    _, v1, _ = two_versions
    a = svc.create_assessment(v1, "First")
    answer(a, {"A1": "compliant", "A2": "partial", "A3": "non_compliant"})
    k = svc.compute_kpis(a)
    assert k["counts"]["not_assessed"] == 1
    assert k["score"] == 50.0  # (1 + 0.5) / 3
    with pytest.raises(svc.AssessmentError):
        svc.complete(a)
    answer(a, {"A4": "not_applicable"})
    svc.complete(a)
    assert a.is_completed
    with pytest.raises(svc.AssessmentError):
        answer(a, {"A1": "partial"})


def test_follow_up_across_versions(two_versions):
    _, v1, v2 = two_versions
    a1 = svc.create_assessment(v1, "2024", assessment_date=date(2024, 1, 1))
    answer(a1, {"A1": "compliant", "A2": "non_compliant", "A3": "partial", "A4": "non_compliant"})
    a1_results = {r.requirement.code: r for r in a1.results}
    a1_results["A3"].comment = "inventory incomplete"
    a1.actions.append(Action(title="Fix roles", requirement=a1_results["A2"].requirement))
    a1.actions.append(Action(title="Old pwd action", requirement=a1_results["A4"].requirement))
    a1.actions.append(Action(title="Done already", status="done"))
    svc.complete(a1)
    db.session.flush()

    a2 = svc.create_assessment(v2, "2025", assessment_date=date(2025, 1, 1), based_on=a1, carry_actions=True)
    res = {r.requirement.code: r for r in a2.results}
    assert res["A1"].status == "compliant" and not res["A1"].needs_review
    assert res["A2"].status == "non_compliant" and res["A2"].needs_review  # reworded
    assert res["B3"].comment == "inventory incomplete"  # followed the renumbering
    assert res["B5"].status is None  # new requirement

    carried = {act.title: act for act in a2.actions}
    assert set(carried) == {"Fix roles", "Old pwd action"}
    assert carried["Fix roles"].requirement.code == "A2"
    assert carried["Old pwd action"].requirement is None  # requirement removed in v2
    assert {act.title: act.status for act in a1.actions}["Fix roles"] == "carried_over"

    answer(a2, {"A2": "compliant", "B3": "non_compliant", "B5": "partial"})
    svc.complete(a2)
    assert svc.previous_assessment(a2) is a1
    diff = svc.compare(a2, a1)
    codes = {k: sorted((c or o).requirement.code for c, o in v) for k, v in diff.items()}
    assert codes == {"improved": ["A2"], "regressed": ["B3"], "changed": [], "unchanged": ["A1"],
                     "new": ["B5"], "removed": ["A4"]}

    evo = svc.evolution(a2.framework)
    assert [p["assessment"].name for p in evo["points"]] == ["2024", "2025"]
    matrix = {row["requirement"].code: row["cells"] for row in evo["matrix"]}
    assert matrix["B3"] == ["partial", "non_compliant"]
    assert matrix["B5"] == [None, "partial"]
    assert matrix["A4"] == ["non_compliant", None]


def test_overdue_action(app):
    act = Action(title="t", status="open", due_date=date.today() - timedelta(days=1))
    assert act.is_overdue
    act.set_status("done")
    assert not act.is_overdue and act.closed_at is not None
    act.set_status("open")
    assert act.closed_at is None


# --- HTTP --------------------------------------------------------------------

def test_full_http_flow(client):
    client.post("/repositories/new", data={"name": "Corp"})
    fw = Framework.query.one()
    xlsx = build_template()
    resp = client.post(f"/repositories/{fw.id}/versions/new",
                       data={"label": "v1", "mode": "upload", "file": (xlsx, "reqs.xlsx")},
                       content_type="multipart/form-data", follow_redirects=True)
    assert b"Draft version v1 created with 3" in resp.data
    v = fw.versions[0]
    client.post(f"/repositories/versions/{v.id}/publish")
    assert v.is_published

    resp = client.post("/assessments/new", data={"version_id": v.id, "name": "Q1"})
    a = Assessment.query.one()
    assert resp.status_code == 302

    r0, r1, r2 = svc.sorted_results(a)
    resp = client.post(f"/assessments/{a.id}/results/{r0.id}", json={"status": "compliant", "comment": "ok"})
    assert resp.json["counts"]["compliant"] == 1
    assert client.post(f"/assessments/{a.id}/results/{r0.id}", json={"status": "bogus"}).status_code == 400

    # Non-JS full-form save.
    client.post(f"/assessments/{a.id}/save", data={
        f"status_{r1.id}": "partial", f"comment_{r1.id}": "half", f"evidence_{r1.id}": "",
        f"status_{r2.id}": "", f"comment_{r2.id}": "", f"evidence_{r2.id}": ""})
    assert r1.status == "partial" and r1.comment == "half"

    client.post("/actions/new", data={"assessment_id": a.id, "requirement_id": r1.requirement_id,
                                      "title": "Finish it", "due_date": "2030-01-01", "priority": "high"})
    assert Action.query.one().title == "Finish it"
    assert client.post("/actions/new", data={"assessment_id": a.id, "title": ""}).status_code == 200

    client.post(f"/assessments/{a.id}/complete", data={})
    assert not a.is_completed  # one requirement still unanswered
    client.post(f"/assessments/{a.id}/complete", data={"allow_unanswered": "1", "assessment_date": "2026-03-31"})
    assert a.is_completed and a.assessment_date == date(2026, 3, 31)

    for url in ["/", f"/assessments/{a.id}", f"/assessments/{a.id}/report", f"/repositories/{fw.id}/evolution",
                "/actions/", f"/repositories/versions/{v.id}"]:
        assert client.get(url).status_code == 200, url

    wb = load_workbook(io.BytesIO(client.get(f"/assessments/{a.id}/export.xlsx").data))
    assert wb.sheetnames == ["Summary", "Results", "Action plan"]
    assert wb["Action plan"]["B2"].value == "Finish it"


def test_redirects_stay_local(client):
    client.post("/repositories/new", data={"name": "Corp"})
    resp = client.post("/actions/1/delete", headers={"Referer": "https://evil.example/x"})
    assert resp.status_code == 404
    from app.routes.utils import safe_next
    assert safe_next("https://evil.example/x?a=1", "/d") == "/x?a=1"
    assert safe_next("//evil.example", "/d") == "/d"
