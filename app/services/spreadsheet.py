"""Spreadsheet template, import of requirement lists, and Excel exports."""

import csv
import io
from dataclasses import dataclass

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import ACTION_LABELS, RESULT_LABELS

# (field, header, required, width, help)
COLUMNS = [
    ("code", "Code", True, 14, "Unique identifier of the requirement in this version (e.g. A.5.1). Keep as text."),
    ("section", "Section", False, 30, "Chapter / domain / theme used to group requirements."),
    ("title", "Title", True, 50, "Short statement of the requirement."),
    ("description", "Description", False, 70, "Full text of the requirement."),
    ("guidance", "Guidance", False, 50, "Optional implementation or audit guidance."),
    (
        "previous_code",
        "Previous code",
        False,
        16,
        "Only for a new version whose code changed: the code of the same requirement in the "
        "previous version. Leave empty when the code is unchanged or the requirement is new.",
    ),
]
FIELDS = [c[0] for c in COLUMNS]
REQUIRED = [c[0] for c in COLUMNS if c[2]]

HEADER_ALIASES = {
    "code": {"code", "id", "ref", "reference", "requirement id", "control", "control id"},
    "section": {"section", "chapter", "domain", "category", "theme", "clause"},
    "title": {"title", "name", "requirement", "control name", "summary"},
    "description": {"description", "text", "details", "requirement text"},
    "guidance": {"guidance", "guidelines", "notes", "implementation guidance"},
    "previous_code": {"previous code", "previous id", "old code", "former code", "previous_code"},
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


@dataclass
class RequirementRow:
    code: str
    title: str
    section: str = ""
    description: str = ""
    guidance: str = ""
    previous_code: str = ""
    line: int = 0


class SpreadsheetError(Exception):
    def __init__(self, errors):
        super().__init__("; ".join(errors))
        self.errors = errors


def _cell_to_str(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _style_header(ws, widths):
    for idx, width in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"


def _to_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_template(requirements=None):
    """Return an .xlsx (BytesIO) in the import format, optionally pre-filled."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Requirements"
    ws.append([c[1] for c in COLUMNS])
    _style_header(ws, [c[3] for c in COLUMNS])
    ws.column_dimensions["A"].number_format = "@"

    rows = requirements
    if rows is None:
        rows = [
            RequirementRow("1.1", "Define an information security policy", "1. Governance",
                           "A policy shall be defined, approved by management and communicated.",
                           "Review at least yearly."),
            RequirementRow("1.2", "Assign security roles and responsibilities", "1. Governance",
                           "Roles and responsibilities shall be defined and allocated."),
            RequirementRow("2.1", "Maintain an asset inventory", "2. Asset management",
                           "An inventory of information assets and owners shall be maintained."),
        ]
    for r in rows:
        ws.append([r.code, r.section, r.title, r.description, r.guidance, getattr(r, "previous_code", "") or ""])
        for col in range(1, len(COLUMNS) + 1):
            ws.cell(row=ws.max_row, column=col).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=ws.max_row, column=1).number_format = "@"

    help_ws = wb.create_sheet("Instructions")
    help_ws.append(["Column", "Required", "Meaning"])
    _style_header(help_ws, [18, 10, 100])
    for _, header, required, _, text in COLUMNS:
        help_ws.append([header, "yes" if required else "no", text])
    help_ws.append([])
    for line in [
        "One row per requirement, in the 'Requirements' sheet (or the first sheet). Row order is kept.",
        "Header names are matched case-insensitively; extra columns are ignored.",
        "New version of an existing repository: requirements are matched to the previous version by Code "
        "(or by 'Previous code' when filled). The tool then flags each one as new, modified or unchanged "
        "and lists the removed ones.",
        "Tip: export the current version from the tool, edit it, and upload it as the next version.",
    ]:
        help_ws.append(["", "", line])
    for row in help_ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return _to_bytes(wb)


def _read_rows(filename, stream):
    name = (filename or "").lower()
    if name.endswith(".csv"):
        text = stream.read().decode("utf-8-sig")
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t") if text.strip() else csv.excel
        return list(csv.reader(io.StringIO(text), dialect))
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(stream, read_only=True, data_only=True)
        ws = wb["Requirements"] if "Requirements" in wb.sheetnames else wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    raise SpreadsheetError(["Unsupported file type: please upload an .xlsx or .csv file."])


def parse_requirements(filename, stream):
    """Parse an uploaded file into RequirementRow objects. Raises SpreadsheetError."""
    raw = _read_rows(filename, stream)
    if not raw:
        raise SpreadsheetError(["The file is empty."])

    header = [_cell_to_str(h).lower() for h in raw[0]]
    mapping = {}
    for idx, h in enumerate(header):
        for field, aliases in HEADER_ALIASES.items():
            if h in aliases and field not in mapping:
                mapping[field] = idx
    missing = [f for f in REQUIRED if f not in mapping]
    if missing:
        headers = ", ".join(dict((c[0], c[1]) for c in COLUMNS)[f] for f in missing)
        raise SpreadsheetError([f"Missing required column(s): {headers}. Download the template for the expected format."])

    rows, errors, seen = [], [], {}
    for line_no, values in enumerate(raw[1:], start=2):
        values = list(values) + [None] * (len(header) - len(values))
        data = {f: _cell_to_str(values[i]) for f, i in mapping.items()}
        if not any(data.values()):
            continue  # skip blank lines
        if not data.get("code"):
            errors.append(f"Line {line_no}: missing Code.")
            continue
        if not data.get("title"):
            errors.append(f"Line {line_no}: missing Title for code '{data['code']}'.")
            continue
        if data["code"] in seen:
            errors.append(f"Line {line_no}: duplicate code '{data['code']}' (already on line {seen[data['code']]}).")
            continue
        seen[data["code"]] = line_no
        rows.append(RequirementRow(line=line_no, **{f: data.get(f, "") for f in FIELDS}))

    if errors:
        raise SpreadsheetError(errors)
    if not rows:
        raise SpreadsheetError(["No requirement found in the file."])
    return rows


def export_assessment(assessment, kpis, results):
    """Excel report of an assessment: summary, results and action plan."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 60
    summary = [
        ("Assessment", assessment.name),
        ("Repository", assessment.framework.name),
        ("Version", assessment.version.label),
        ("Status", "Completed" if assessment.is_completed else "In progress"),
        ("Assessment date", assessment.reference_date.isoformat()),
        ("Assessor", assessment.assessor),
        ("Scope", assessment.scope),
        ("General comment", assessment.general_comment),
        ("", ""),
        ("Requirements", kpis["total"]),
    ]
    summary += [(RESULT_LABELS[k], kpis["counts"][k]) for k in RESULT_LABELS]
    summary += [
        ("Not assessed", kpis["counts"]["not_assessed"]),
        ("Compliance score (%)", kpis["score"] if kpis["score"] is not None else "n/a"),
        ("Open actions", kpis["actions"]["open_total"]),
        ("Overdue actions", kpis["actions"]["overdue"]),
    ]
    for label, value in summary:
        ws.append([label, value])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)

    rs = wb.create_sheet("Results")
    rs.append(["Code", "Section", "Title", "Status", "Comment", "Evidence", "Change vs previous version"])
    _style_header(rs, [14, 28, 50, 20, 50, 40, 18])
    for r in results:
        req = r.requirement
        rs.append([req.code, req.section, req.title, RESULT_LABELS.get(r.status, "Not assessed"),
                   r.comment, r.evidence, req.change_type or ""])

    acts = wb.create_sheet("Action plan")
    acts.append(["Requirement", "Action", "Description", "Owner", "Priority", "Due date", "Status", "Closed"])
    _style_header(acts, [14, 40, 50, 20, 10, 12, 14, 12])
    for a in assessment.actions:
        acts.append([a.requirement.code if a.requirement else "", a.title, a.description, a.owner, a.priority,
                     a.due_date, ACTION_LABELS[a.status], a.closed_at.date() if a.closed_at else None])
    return _to_bytes(wb)
